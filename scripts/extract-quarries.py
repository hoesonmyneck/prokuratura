"""
Разбор ведомости резервных месторождений в GeoJSON.

Исходник — `carrier.xlsx`, «Сведения по резервным месторождениям
Кызылординской области». Каждое месторождение занимает несколько строк:
одна строка — одна угловая точка контура, координаты в градусах, минутах
и секундах, разложенные по отдельным столбцам.

Столбцы 11-15 (запасы, протокол, единица измерения) объединены по вертикали
на всю группу, поэтому значение стоит только в первой строке — переносим его
на всё месторождение.

Одна запись может содержать НЕСКОЛЬКО площадок. Иногда это видно по нумерации
(у №20 вершины идут «1 2 3 4 1 2 3 4»), но не всегда: у №21 номера сквозные,
с 1 по 14, а площадок три — они разнесены на 6 и 14 км. Поэтому вершины
разбиваются по расстоянию, а не по номерам.

В ведомости есть опечатки в координатах: у №14 у одной вершины широта 42°
вместо 45° (полигон уходил на 333 км в Узбекистан), у №8 минуты 54' вместо 34'.
Исправляются они только там, где это доказуемо: если подстановка ОДНОГО
разряда из соседних вершин той же площадки ставит точку на место, правка
применяется, иначе вершина отбрасывается. Каждая правка печатается.

Требуется: pip install openpyxl
Запуск:    python scripts/extract-quarries.py
Результат: src/data/quarries.json
"""

import json
import re
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib.deposits import build_rings, dms, ring_area_ha  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "carrier.xlsx"
OUT = ROOT / "src" / "data" / "quarries.json"

FIRST_ROW = 11
# Столбцы (нумерация с единицы): № | наименование | вид ПИ | № точки |
# долгота г/м/с | широта г/м/с | запасы А+В | А+В+С1 | С2 | протокол | единица
COL = {
    "num": 1, "name": 2, "kind": 3, "point": 4,
    "lon": (5, 6, 7), "lat": (8, 9, 10),
    "ab": 11, "abc1": 12, "c2": 13, "protocol": 14, "unit": 15,
}


# Вид полезного ископаемого записан вразнобой: «супеси» и «супесь», «пгс»
# и «ПГС». Сводим к одному написанию для группировки, исходное оставляем как есть.
KIND_ALIASES = {
    "супеси": "супесь",
    "супесь": "супесь",
    "суглинки": "суглинок",
    "суглинок": "суглинок",
    "пгс": "ПГС",
    "песчано-гравийная смесь": "ПГС",
}


def normal_kind(value):
    if not value:
        return None
    key = value.strip().lower()
    return KIND_ALIASES.get(key, value.strip().lower())


def number(value):
    if value is None:
        return None
    try:
        return float(str(value).strip().replace(",", "."))
    except ValueError:
        return None


def clean(value):
    """Пробелы и переносы внутри ячеек — обычное дело, схлопываем."""
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Не найден {SRC.name} в корне проекта")

    sheet = openpyxl.load_workbook(SRC, data_only=True).active
    groups = []
    current = None
    kind = None

    for row_index, row in enumerate(
        sheet.iter_rows(min_row=FIRST_ROW, max_row=sheet.max_row, values_only=True), start=FIRST_ROW
    ):
        cell = [row[i - 1] if i - 1 < len(row) else None for i in range(1, 16)]
        if clean(cell[COL["kind"] - 1]):
            # Вид ПИ объединён на всю группу — держим последний увиденный.
            kind = clean(cell[COL["kind"] - 1])

        parts = [number(cell[i - 1]) for i in (*COL["lon"], *COL["lat"])]
        has_point = all(p is not None for p in parts)
        label = clean(cell[COL["num"] - 1])
        name = clean(cell[COL["name"] - 1])

        if label and (name or has_point):
            current = {
                "num": label,
                "name": name,
                "kind": kind,
                "ab": clean(cell[COL["ab"] - 1]),
                "abc1": clean(cell[COL["abc1"] - 1]),
                "c2": clean(cell[COL["c2"] - 1]),
                "protocol": clean(cell[COL["protocol"] - 1]),
                "unit": clean(cell[COL["unit"] - 1]),
                "points": [],
            }
            groups.append(current)

        if has_point and current is not None:
            current["points"].append({
                "row": row_index,
                "order": len(current["points"]),
                "seq": number(cell[COL["point"] - 1]),
                "raw": parts,
                # В этой ведомости первой идёт долгота.
                "build": lambda raw: (dms(*raw[:3]), dms(*raw[3:])),
                "xy": (dms(*parts[:3]), dms(*parts[3:])),
            })

    features = []
    split = 0
    repaired = []
    discarded = []
    review = []
    for index, g in enumerate(groups, start=1):
        rings = build_rings(
            g["points"], f"№{g['num']}", ("долготы", "широты"), review, repaired, discarded, print
        )
        if len(rings) > 1:
            split += 1

        if not rings:
            continue

        area = round(sum(ring_area_ha(r[0][:-1]) for r in rings), 1)
        features.append({
            "type": "Feature",
            "id": index,
            "properties": {
                "num": g["num"],
                "name": g["name"],
                "kind": g["kind"],
                "kindNorm": normal_kind(g["kind"]),
                "reservesAB": g["ab"],
                "reservesABC1": g["abc1"],
                "reservesC2": g["c2"],
                "unit": g["unit"],
                "protocol": g["protocol"],
                "areaHa": area,
                "plots": len(rings),
            },
            "geometry": {"type": "MultiPolygon", "coordinates": rings},
        })

    OUT.write_text(
        json.dumps({"type": "FeatureCollection", "features": features}, ensure_ascii=False),
        encoding="utf-8",
    )

    kinds = {}
    for f in features:
        k = f["properties"]["kindNorm"] or "не указан"
        kinds[k] = kinds.get(k, 0) + 1
    total = sum(f["properties"]["areaHa"] for f in features)
    print(f"Месторождений: {len(features)}, площадок: {sum(f['properties']['plots'] for f in features)}")
    print(f"Записей с несколькими площадками: {split}")
    print(f"Исправлено опечаток в координатах: {len(repaired)}; отброшено вершин: {len(discarded)}")
    if review:
        print("Шаги вблизи порога разреза — стоит проверить глазами:")
        for line in review:
            print(line)
    print(f"Суммарная площадь: {total:,.0f} га".replace(",", " "))
    print("По виду полезного ископаемого:")
    for k, n in sorted(kinds.items(), key=lambda kv: -kv[1]):
        print(f"  {k:<28} {n}")
    print(f"\nЗаписано: {OUT}")


if __name__ == "__main__":
    main()
